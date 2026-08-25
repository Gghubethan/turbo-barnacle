#!/usr/bin/env python3
"""
ats_check.py — ATS（简历筛选系统）兼容性扫描（流程 ⑦）。

职责：只查**机器可判定**的项，给出 🔴致命 / 🟡高风险 / ⚪建议 三级结论：
  文本能否提取、分栏、文本框、表格布局、图片、页眉页脚藏联系方式、
  字体是否常见、项目符号是否会变乱码、联系方式能否被正则抓到、
  日期格式是否一致、模块标题是否用标准名、页数是否超标、文件名。

用法：
  python3 ats_check.py 简历.docx
  python3 ats_check.py 简历.pdf --json
  python3 ats_check.py resume-work/resume.json     # 只查内容层

判断不了的（关键词匹配、缩写展开、内容取舍）见 references/07-ats-check.md 的人工判断项。
不打分——"ATS 通过率 87%"这类数字都是编的。
依赖：python-docx（.docx）、pdfminer.six / pypdf（.pdf）。
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from parse_resume import (  # noqa: E402
    SECTION_KEYWORDS, RE_EMAIL, RE_PHONE, match_section, read_pdf,
)

FATAL, HIGH, LOW = "🔴", "🟡", "⚪"

SAFE_FONTS = {
    "calibri", "arial", "helvetica", "times new roman", "georgia", "verdana", "tahoma",
    "segoe ui", "cambria", "garamond", "book antiqua", "palatino linotype", "roboto",
    "微软雅黑", "microsoft yahei", "宋体", "simsun", "黑体", "simhei", "等线", "dengxian",
    "楷体", "kaiti", "仿宋", "fangsong", "思源黑体", "source han sans sc", "source han sans",
    "pingfang sc", "noto sans cjk sc", "noto sans sc", "yu gothic",
}
RISKY_BULLETS = set("❖✦➢➤◆◇▪‣∙※★☆♦→▶")
DATE_STYLES = {
    "2021.03": re.compile(r"(19|20)\d{2}\.\d{1,2}"),
    "2021-03": re.compile(r"(19|20)\d{2}-\d{1,2}(?!\d)"),
    "2021/03": re.compile(r"(19|20)\d{2}/\d{1,2}"),
    "2021年3月": re.compile(r"(19|20)\d{2}\s*年\s*\d{1,2}\s*月"),
    "03/2021": re.compile(r"(?<!\d)\d{1,2}/(19|20)\d{2}"),
    "Mar 2021": re.compile(r"(?i)\b(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?\s+(19|20)\d{2}"),
}


def finding(level: str, item: str, detail: str, fix: str) -> dict:
    return {"level": level, "item": item, "detail": detail, "fix": fix}


# ---- 内容层（三种输入都查） ------------------------------------------------
def check_text(text: str, headings: list[str]) -> list[dict]:
    out: list[dict] = []
    if len(text.strip()) < 80:
        out.append(finding(FATAL, "文本提取", "几乎提取不到文字（扫描件或图片版简历）",
                           "ATS 完全读不到内容 = 直接出局。请提供 DOCX 或文字版 PDF"))
        return out

    if not RE_EMAIL.search(text):
        out.append(finding(FATAL, "邮箱", "正文里没抓到可识别的邮箱",
                           "把邮箱写成标准格式放在正文首行，别拆成两段、别放页眉、别写成图片"))
    if not RE_PHONE.search(text):
        out.append(finding(FATAL, "电话", "正文里没抓到可识别的手机号",
                           "写成 138-1234-5678 或 13812345678，中间别插空格或全角符号"))

    styles = [name for name, pattern in DATE_STYLES.items() if pattern.search(text)]
    if len(styles) > 1:
        out.append(finding(HIGH, "日期格式", f"混用了 {len(styles)} 种写法：{'、'.join(styles)}",
                           "全篇统一一种（推荐 2021.03 - 2023.06），否则时间线会被解析错"))

    if headings:
        standard = {kw for _, _, kws in SECTION_KEYWORDS for kw in kws}
        odd = [h for h in headings
               if not any(kw.replace(" ", "").lower() in re.sub(r"\s", "", h).lower()
                          for kw in standard)]
        if odd:
            out.append(finding(HIGH, "模块标题", f"非标准标题：{'、'.join(odd[:4])}",
                               "ATS 按标题分段。请用「工作经历/教育背景/项目经历/技能」这类标准名"))
    else:
        out.append(finding(HIGH, "模块标题", "没识别到任何标准模块标题",
                           "加上标准模块标题，否则 ATS 无法把内容归类"))

    risky = RISKY_BULLETS & set(text)
    if risky:
        out.append(finding(HIGH, "项目符号", f"使用了特殊符号：{' '.join(sorted(risky))}",
                           "统一换成 • ，特殊符号转纯文本时可能变成乱码或问号"))
    return out


# ---- DOCX ------------------------------------------------------------------
def check_docx(path: Path) -> tuple[list[dict], str]:
    try:
        import docx
        from docx.oxml.ns import qn
    except ImportError as exc:
        sys.exit(f"依赖导入失败（{exc}）：pip install -r .claude/skills/resume-expert/requirements.txt")

    doc = docx.Document(str(path))
    out: list[dict] = []
    body_text = "\n".join(p.text for p in doc.paragraphs)
    table_text = "\n".join(cell.text for t in doc.tables for row in t.rows for cell in row.cells)
    text = body_text + "\n" + table_text
    xml = doc.element.body.xml

    for section in doc.sections:
        cols = section._sectPr.find(qn("w:cols"))
        if cols is not None and (cols.get(qn("w:num")) or "1") != "1":
            out.append(finding(FATAL, "分栏", f"检测到 {cols.get(qn('w:num'))} 栏排版",
                               "改成单栏。双栏会让解析器把左右两栏文字交错串行，读出来是乱的"))
            break

    if "txbxContent" in xml:
        out.append(finding(FATAL, "文本框", "文档里有文本框",
                           "把文本框内容移到正文段落，多数解析器直接丢弃文本框"))

    if doc.tables:
        share = len(table_text.strip()) / max(len(text.strip()), 1)
        level = FATAL if share > 0.3 else LOW
        out.append(finding(level, "表格", f"{len(doc.tables)} 个表格，承载约 {share:.0%} 的文字",
                           "表格用于版面时解析风险高；用 build_docx.py 重建为单栏，或只保留纯列表"))

    if "<w:drawing" in xml or "<w:pict" in xml:
        out.append(finding(HIGH, "图片/图形", "文档含图片或图形",
                           "确认没有把文字信息（联系方式、技能条、图表）画进图里——图里的字等于没写"))

    for section in doc.sections:
        for part, label in ((section.header, "页眉"), (section.footer, "页脚")):
            part_text = "\n".join(p.text for p in part.paragraphs).strip()
            if part_text and (RE_EMAIL.search(part_text) or RE_PHONE.search(part_text)):
                out.append(finding(FATAL, label, f"{label}里放了联系方式",
                                   f"移到正文首行——多数 ATS 不读{label}，等于没留联系方式"))

    fonts = {f.get(qn("w:ascii")) for f in doc.element.body.iter(qn("w:rFonts"))}
    fonts |= {f.get(qn("w:eastAsia")) for f in doc.element.body.iter(qn("w:rFonts"))}
    unsafe = sorted({f for f in fonts if f and f.lower() not in SAFE_FONTS})
    if unsafe:
        out.append(finding(LOW, "字体", f"非常见字体：{'、'.join(unsafe[:4])}",
                           "换成雅黑/宋体/Calibri/Arial 等常见字体，冷门字体在对方电脑上会回退、排版全乱"))

    est_pages = max(1, round(len(re.sub(r"\s", "", text)) / 1600 + 0.4))
    if est_pages > 2:
        out.append(finding(HIGH, "篇幅", f"按字数估算约 {est_pages} 页",
                           "压到 2 页以内，删减顺序见 references/06-layout-review.md"))

    headings = [p.text.strip() for p in doc.paragraphs
                if p.text.strip() and match_section(p.text.strip(),
                                                    p.style.name if p.style is not None else "")]
    out.extend(check_text(text, headings))
    return out, text


# ---- PDF -------------------------------------------------------------------
def check_pdf(path: Path) -> tuple[list[dict], str]:
    lines, warnings = read_pdf(path)
    text = "\n".join(ln["text"] for ln in lines)
    out: list[dict] = []
    for warning in warnings:
        if warning.startswith("no_text_extracted"):
            out.append(finding(FATAL, "文本提取", "PDF 里提取不到文字（扫描件或转成了图片）",
                               "重新从 Word 导出为 PDF（另存为 PDF，不要「打印成图片」），或直接投 DOCX"))
        elif warning.startswith("multi_column_suspected"):
            out.append(finding(FATAL, "分栏", "疑似双栏排版，文字左右交错",
                               "改成单栏后重新导出"))

    try:
        import pypdf

        pages = len(pypdf.PdfReader(str(path)).pages)
        if pages > 2:
            out.append(finding(HIGH, "篇幅", f"共 {pages} 页", "压到 2 页以内"))
    except Exception:  # noqa: BLE001 — 数不出页数不阻断其他检查
        pass

    headings = [ln["text"] for ln in lines if match_section(ln["text"], "")]
    out.extend(check_text(text, headings))
    return out, text


def check_json(path: Path) -> tuple[list[dict], str]:
    resume = json.loads(path.read_text(encoding="utf-8"))
    text = "\n".join(resume.get("raw_lines", []))
    headings = [s.get("title", "") for s in resume.get("sections", [])]
    out = check_text(text, headings)
    basics = resume.get("basics", {})
    if basics.get("email") and not RE_EMAIL.search(text):
        text += "\n" + basics["email"]
    return out, text


def check_filename(path: Path) -> list[dict]:
    name = path.stem
    if re.search(r"(最终版|最新版|副本|final|copy|\(\d\)|_\d{6,})", name, re.IGNORECASE):
        return [finding(LOW, "文件名", f"「{name}」看起来像工作文件",
                        "改成「姓名-目标岗位-简历」，HR 下载后一眼能认出是谁")]
    if not re.search(r"(简历|resume|cv)", name, re.IGNORECASE):
        return [finding(LOW, "文件名", f"「{name}」不含「简历」字样",
                        "建议「姓名-目标岗位-简历.docx」")]
    return []


def main() -> int:
    ap = argparse.ArgumentParser(description="ATS 兼容性扫描（流程 ⑦）")
    ap.add_argument("input", type=Path, help=".docx / .pdf / resume.json")
    ap.add_argument("--json", action="store_true", help="输出 JSON，便于程序处理")
    args = ap.parse_args()

    if not args.input.exists():
        sys.exit(f"文件不存在：{args.input}")
    suffix = args.input.suffix.lower()
    if suffix == ".docx":
        findings, _ = check_docx(args.input)
    elif suffix == ".pdf":
        findings, _ = check_pdf(args.input)
    elif suffix == ".json":
        findings, _ = check_json(args.input)
    else:
        sys.exit(f"不支持的格式：{suffix}（支持 .docx / .pdf / resume.json）")
    findings += check_filename(args.input)

    order = {FATAL: 0, HIGH: 1, LOW: 2}
    findings.sort(key=lambda f: order[f["level"]])
    counts = {level: sum(1 for f in findings if f["level"] == level) for level in (FATAL, HIGH, LOW)}

    if args.json:
        print(json.dumps({"file": str(args.input), "counts": counts, "findings": findings},
                         ensure_ascii=False, indent=2))
        return 0

    print(f"ATS 扫描：{args.input.name}")
    print(f"致命 {counts[FATAL]} 项 ｜ 高风险 {counts[HIGH]} 项 ｜ 建议 {counts[LOW]} 项\n")
    for f in findings:
        print(f"{f['level']} {f['item']}：{f['detail']}\n   → {f['fix']}")
    if not findings:
        print("未发现机器可判定的问题。")
    print("\n结论：" + ("有致命项，修复后再投递。" if counts[FATAL]
                      else "可以投递；建议项按需处理。"))
    print("关键词匹配、缩写展开等人工判断项见 references/07-ats-check.md。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
