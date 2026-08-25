#!/usr/bin/env python3
"""
format_docx.py — 就地规范一份现有 DOCX 的排版，不重建、不改文字（流程 ⑧ 的保守路线）。

职责：
1) 按 assets/format-profiles.json 统一页边距、中西文字体、字号、行距、段间距、左对齐；
2) 识别姓名 / 联系方式 / 模块标题 / 条目标题 / bullet / 正文，分别套用规格；
3) 规范项目符号（统一符号 + 悬挂缩进），删除用来凑间距的空段落；
4) 强制单栏；对文本框、表格布局、页眉页脚里的联系方式等**只报告不擅自改结构**。

用法：
  python3 format_docx.py 原简历.docx -o 排版修正版.docx
  python3 format_docx.py 原简历.docx --dry-run          # 只列出会改什么，不写文件

文字内容一字不动（仅项目符号字符会被统一）；永不覆盖输入文件。
适用于"用户想保留自己的版式"；已走完改写流程的，用 build_docx.py 重建更干净。
依赖：python-docx。
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

try:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
except ImportError as exc:
    sys.exit(f"依赖导入失败（{exc}）：pip install -r .claude/skills/resume-expert/requirements.txt")

sys.path.insert(0, str(Path(__file__).resolve().parent))
from parse_resume import RE_DATE_RANGE, RE_EMAIL, RE_PHONE, is_bullet, match_section, strip_bullet  # noqa: E402

PROFILES = Path(__file__).resolve().parent.parent / "assets" / "format-profiles.json"


def load_profile(name: str | None, doc) -> dict:
    data = json.loads(PROFILES.read_text(encoding="utf-8"))
    if name:
        return data[name]
    text = "\n".join(p.text for p in doc.paragraphs)
    cjk = len(re.findall(r"[一-龥]", text))
    return data["zh"] if cjk > max(len(text), 1) * 0.1 else data["en"]


def classify(paragraphs, index: int, para) -> str:
    """判断段落角色。顺序敏感：姓名和联系方式只可能在开头。"""
    text = para.text.strip()
    if not text:
        return "empty"
    style = para.style.name if para.style is not None else ""
    if index <= 3:
        if RE_EMAIL.search(text) or RE_PHONE.search(text) or text.count("|") >= 2:
            return "contact"
        if index == 0 and len(text) <= 12:
            return "name"
    if match_section(text, style):
        return "heading"
    if is_bullet(text, style):
        return "bullet"
    if RE_DATE_RANGE.search(text):
        return "entry_title"
    return "body"


def style_run(run, profile: dict, size: float, bold: bool | None) -> None:
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    run.font.name = profile["fonts"]["ascii"]
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), profile["fonts"]["ascii"])
    rfonts.set(qn("w:hAnsi"), profile["fonts"]["ascii"])
    rfonts.set(qn("w:eastAsia"), profile["fonts"]["east_asia"])


def apply_para(para, role: str, profile: dict) -> str:
    """套用某个角色的排版规格，返回一句人类可读的改动说明。"""
    sizes, spacing, bullet = profile["sizes"], profile["spacing"], profile["bullet"]
    spec = {
        "name": (sizes["name"], True, 0, spacing["name_after_pt"], False),
        "contact": (sizes["contact"], None, 0, spacing["contact_after_pt"], False),
        "heading": (sizes["section_heading"], True, spacing["section_before_pt"],
                    spacing["section_after_pt"], True),
        "entry_title": (sizes["entry_title"], None, spacing["entry_before_pt"],
                        spacing["entry_after_pt"], True),
        "bullet": (sizes["bullet"], None, 0, spacing["bullet_after_pt"], False),
        "body": (sizes["body"], None, 0, spacing["bullet_after_pt"], False),
    }[role]
    size, bold, before, after, keep_next = spec

    fmt = para.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.LEFT          # 禁止两端对齐
    fmt.line_spacing = profile["spacing"]["line_spacing"]
    fmt.space_before, fmt.space_after = Pt(before), Pt(after)
    fmt.keep_with_next = keep_next
    fmt.keep_together = profile["rules"]["keep_entry_together"]

    note = ""
    if role == "bullet":
        fmt.left_indent = Cm(bullet["indent_left_cm"])
        fmt.first_line_indent = Cm(-bullet["hanging_cm"])
        raw = para.text.strip()
        stripped = strip_bullet(raw)
        marker = raw[: len(raw) - len(stripped)].strip()
        if marker and marker != bullet["char"] and not marker[0].isdigit():
            for run in para.runs:                     # 统一项目符号字符
                if marker[0] in run.text:
                    run.text = run.text.replace(marker, bullet["char"], 1)
                    note = f"符号 {marker} → {bullet['char']}"
                    break
    else:
        fmt.left_indent = Cm(0)
        fmt.first_line_indent = Cm(0)

    for run in para.runs:
        style_run(run, profile, size, bold)
    return note


def delete_para(para) -> None:
    para._element.getparent().remove(para._element)


def structural_warnings(doc) -> list[str]:
    """只报告不擅自改：这些问题需要重建版式，交给 build_docx.py 或人工决定。"""
    warnings: list[str] = []
    xml = doc.element.body.xml
    if "txbxContent" in xml:
        warnings.append("🔴 存在文本框：多数 ATS 解析器会直接丢弃文本框里的文字")
    if doc.tables:
        warnings.append(f"🟡 存在 {len(doc.tables)} 个表格：若用来做版面，建议改用 build_docx.py 重建为单栏")
    if "<w:drawing" in xml or "<w:pict" in xml:
        warnings.append("🟡 存在图片/图形：图里的文字等于没写")
    for section in doc.sections:
        for part, label in ((section.header, "页眉"), (section.footer, "页脚")):
            text = "\n".join(p.text for p in part.paragraphs).strip()
            if text and (RE_EMAIL.search(text) or RE_PHONE.search(text)):
                warnings.append(f"🔴 {label}里有联系方式：ATS 通常不读{label}，请移到正文首行")
            elif text:
                warnings.append(f"⚪ {label}有内容：{text[:20]}")
    return warnings


def process(doc, profile: dict) -> tuple[list[str], int]:
    changes: list[str] = []
    page = profile["page"]
    for section in doc.sections:
        section.top_margin, section.bottom_margin = Cm(page["margin_top_cm"]), Cm(page["margin_bottom_cm"])
        section.left_margin, section.right_margin = Cm(page["margin_left_cm"]), Cm(page["margin_right_cm"])
        cols = section._sectPr.find(qn("w:cols"))     # 强制单栏
        if cols is not None and cols.get(qn("w:num")) not in (None, "1"):
            cols.set(qn("w:num"), "1")
            changes.append("分栏 → 单栏（双栏会让 ATS 把左右文字交错读取）")

    paragraphs = list(doc.paragraphs)
    to_delete = []
    counts: dict[str, int] = {}
    for i, para in enumerate(paragraphs):
        role = classify(paragraphs, i, para)
        if role == "empty":
            to_delete.append(para)                    # 用段间距控制留白，不用空行
            continue
        note = apply_para(para, role, profile)
        counts[role] = counts.get(role, 0) + 1
        if note:
            changes.append(f"第 {i + 1} 段：{note}")

    for table in doc.tables:                          # 表格里的文字也要统一字体字号
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    if para.text.strip():
                        apply_para(para, "body", profile)

    for para in to_delete:
        delete_para(para)

    summary = "、".join(f"{k}×{v}" for k, v in counts.items())
    changes.insert(0, f"统一字体/字号/行距/对齐：{summary}")
    changes.insert(1, f"页边距 → 上下 {page['margin_top_cm']}cm、左右 {page['margin_left_cm']}cm")
    if to_delete:
        changes.insert(2, f"删除 {len(to_delete)} 个用来凑间距的空段落（改用段前段后间距）")
    return changes, len(to_delete)


def main() -> int:
    ap = argparse.ArgumentParser(description="规范现有 DOCX 的排版，不改文字（流程 ⑧）")
    ap.add_argument("input", type=Path, help="原简历 .docx")
    ap.add_argument("-o", "--output", type=Path, help="输出 .docx（--dry-run 时可省略）")
    ap.add_argument("--profile", choices=["zh", "en"], default=None, help="排版规格，默认自动判断")
    ap.add_argument("--dry-run", action="store_true", help="只列出会改什么，不写文件")
    ap.add_argument("--force", action="store_true", help="允许覆盖已存在的输出文件")
    args = ap.parse_args()

    if not args.input.exists():
        sys.exit(f"文件不存在：{args.input}")
    if args.input.suffix.lower() != ".docx":
        sys.exit("只支持 .docx（PDF 无法就地排版，请用 build_docx.py 重建）")
    if not args.dry_run:
        if not args.output:
            sys.exit("需要 -o 输出路径（或加 --dry-run 只看改动）")
        if args.output.resolve() == args.input.resolve():
            sys.exit("输出不能覆盖输入文件，请换个路径")
        if args.output.exists() and not args.force:
            sys.exit(f"输出文件已存在：{args.output}（加 --force 覆盖）")

    doc = docx.Document(str(args.input))
    profile = load_profile(args.profile, doc)
    warnings = structural_warnings(doc)
    changes, _ = process(doc, profile)

    print(f"排版规格：{profile['label']}")
    for line in changes:
        print(f"  · {line}")
    if warnings:
        print("\n需要你决定的结构问题（脚本不擅自改）：")
        for line in warnings:
            print(f"  {line}")

    if args.dry_run:
        print("\n（dry-run，未写文件）")
        return 0

    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(args.output))
    print(f"\n✓ {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
