"""
resume-expert 脚本回归测试。

不碰网络、不读用户真实简历：全部用临时目录里现造的 DOCX 夹具。
重点验证三件事：解析不丢内容、排版规格真的落到 XML 上、零编造检查能抓到编造。

  python3 -m pytest .claude/skills/resume-expert/tests -v
"""
from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

import pytest

docx = pytest.importorskip("docx", reason="需要 python-docx")

import ats_check  # noqa: E402
import build_docx  # noqa: E402
import diff_report  # noqa: E402
import format_docx  # noqa: E402
import parse_resume  # noqa: E402

SCRIPTS = Path(__file__).resolve().parent.parent / "scripts"

GOOD_LINES = [
    ("张三", None),
    ("求职意向：数据分析师", None),
    ("138-1234-5678 | zhangsan@example.com | 北京市", None),
    ("个人简介", "Heading 1"),
    ("三年互联网数据分析经验。", None),
    ("工作经历", "Heading 1"),
    ("字节跳动 | 数据分析师 | 2021.03 - 2023.06", None),
    ("负责用户增长相关的数据分析工作。", "List Bullet"),
    ("搭建用户增长数据看板，覆盖 12 个核心指标。", "List Bullet"),
    ("某某科技 | 数据分析实习生 | 2020.06 - 2020.12", None),
    ("协助完成日报周报的数据整理。", "List Bullet"),
    ("教育背景", "Heading 1"),
    ("北京某大学 | 统计学 学士 | 2016.09 - 2020.06", None),
    ("专业技能", "Heading 1"),
    ("SQL、Python、Tableau", None),
]


@pytest.fixture
def good_docx(tmp_path: Path) -> Path:
    doc = docx.Document()
    for text, style in GOOD_LINES:
        doc.add_paragraph(text, style=style) if style else doc.add_paragraph(text)
    path = tmp_path / "好简历.docx"
    doc.save(str(path))
    return path


@pytest.fixture
def bad_docx(tmp_path: Path) -> Path:
    """排版糟糕的简历：页眉藏联系方式、特殊符号、两端对齐、双倍行距、空行凑间距。"""
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = docx.Document()
    doc.sections[0].header.paragraphs[0].text = "李四 · 13800001111 · lisi@example.com"
    for text in ("李四", "", "", "北京"):
        doc.add_paragraph(text)
    doc.add_paragraph("工作经历")
    doc.add_paragraph("某公司 | 产品经理 | 2019.05 - 2023.08")
    for text in ("❖ 负责产品需求文档撰写与评审推进。",
                 "❖ 参与用户调研，输出调研报告 20 份，支撑版本规划。",
                 "❖ 协调研发与设计资源，跟进版本按期上线。"):
        para = doc.add_paragraph(text)
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        para.paragraph_format.line_spacing = 2.0
    doc.add_paragraph("教育背景")
    doc.add_paragraph("某某大学 | 工商管理 学士 | 2015.09 - 2019.06")
    path = tmp_path / "bad.docx"
    doc.save(str(path))
    return path


@pytest.fixture
def parsed(good_docx: Path) -> dict:
    return parse_resume.build(good_docx)


# ---- ① 解析 ---------------------------------------------------------------
def test_parse_extracts_basics(parsed: dict) -> None:
    basics = parsed["basics"]
    assert basics["name"] == "张三"
    assert basics["phone"] == "138-1234-5678"
    assert basics["email"] == "zhangsan@example.com"
    assert basics["headline"] == "数据分析师"       # "求职意向：xxx"是内容行，不能当模块标题


def test_parse_sections_and_entries(parsed: dict) -> None:
    ids = [s["id"] for s in parsed["sections"]]
    assert ids == ["summary", "experience", "education", "skills"]
    experience = parsed["sections"][1]
    assert len(experience["entries"]) == 2
    first = experience["entries"][0]
    assert first["org"] == "字节跳动" and first["role"] == "数据分析师"
    assert first["date"] == "2021.03 - 2023.06"
    assert [b["text"] for b in first["bullets"]] == [
        "负责用户增长相关的数据分析工作。",
        "搭建用户增长数据看板，覆盖 12 个核心指标。",
    ]


def test_parse_ids_are_addressable(parsed: dict) -> None:
    """id 要能对应"工作经历 → 字节跳动 → 第2条"这样的定位。"""
    bullet = parsed["sections"][1]["entries"][0]["bullets"][1]
    assert bullet["id"] == "experience-1-b2"
    assert parsed["raw_lines"][bullet["line"] - 1] == bullet["text"]


def test_parse_rejects_legacy_doc(tmp_path: Path) -> None:
    legacy = tmp_path / "简历.doc"
    legacy.write_bytes(b"\xd0\xcf\x11\xe0")
    with pytest.raises(SystemExit):
        parse_resume.build(legacy)


# ---- ⑧ 排版：build_docx ----------------------------------------------------
def test_build_docx_roundtrip_keeps_content(parsed: dict, tmp_path: Path) -> None:
    out = tmp_path / "out.docx"
    profile = build_docx.load_profile(None, parsed)
    build_docx.render(parsed, profile).save(str(out))
    again = parse_resume.build(out)
    assert [s["id"] for s in again["sections"]] == [s["id"] for s in parsed["sections"]]
    before = [b["text"] for s in parsed["sections"] for e in s.get("entries", []) for b in e["bullets"]]
    after = [b["text"] for s in again["sections"] for e in s.get("entries", []) for b in e["bullets"]]
    assert before == after


def test_build_docx_applies_format_spec(parsed: dict, tmp_path: Path) -> None:
    from docx.oxml.ns import qn

    out = tmp_path / "out.docx"
    profile = build_docx.load_profile("zh", parsed)
    build_docx.render(parsed, profile).save(str(out))
    doc = docx.Document(str(out))

    assert round(doc.sections[0].left_margin.cm, 2) == profile["page"]["margin_left_cm"]
    name = doc.paragraphs[0]
    assert name.runs[0].font.size.pt == profile["sizes"]["name"]
    rfonts = name.runs[0]._element.find(qn("w:rPr")).find(qn("w:rFonts"))
    assert rfonts.get(qn("w:eastAsia")) == profile["fonts"]["east_asia"]  # 中西文必须分设
    assert rfonts.get(qn("w:ascii")) == profile["fonts"]["ascii"]

    bullets = [p for p in doc.paragraphs if p.text.startswith(profile["bullet"]["char"])]
    assert bullets, "应当渲染出项目符号段落"
    assert round(bullets[0].paragraph_format.first_line_indent.cm, 1) == \
        -profile["bullet"]["hanging_cm"]                                  # 悬挂缩进
    assert bullets[0].paragraph_format.line_spacing == profile["spacing"]["line_spacing"]


def test_build_docx_border_is_schema_valid(parsed: dict, tmp_path: Path) -> None:
    """pBdr 插错位置 Word 会报文件损坏，必须排在 spacing/jc 之前。"""
    from docx.oxml.ns import qn

    out = tmp_path / "out.docx"
    build_docx.render(parsed, build_docx.load_profile("zh", parsed)).save(str(out))
    for para in docx.Document(str(out)).paragraphs:
        ppr = para._element.find(qn("w:pPr"))
        if ppr is None:
            continue
        tags = [child.tag.split("}")[-1] for child in ppr]
        if "pBdr" in tags:
            for later in ("spacing", "ind", "jc", "rPr"):
                if later in tags:
                    assert tags.index("pBdr") < tags.index(later)
            return
    pytest.fail("没有生成模块标题横线")


def test_build_docx_refuses_overwrite(parsed: dict, tmp_path: Path) -> None:
    src = tmp_path / "resume.json"
    src.write_text(json.dumps(parsed, ensure_ascii=False), encoding="utf-8")
    out = tmp_path / "out.docx"
    out.write_text("占位", encoding="utf-8")
    result = subprocess.run([sys.executable, str(SCRIPTS / "build_docx.py"), str(src), "-o", str(out)],
                            capture_output=True, text=True)
    assert result.returncode != 0 and "已存在" in result.stderr


# ---- ⑧ 排版：format_docx ---------------------------------------------------
def test_format_docx_normalizes_without_changing_text(bad_docx: Path, tmp_path: Path) -> None:
    out = tmp_path / "fixed.docx"
    doc = docx.Document(str(bad_docx))
    profile = format_docx.load_profile("zh", doc)
    format_docx.process(doc, profile)
    doc.save(str(out))

    fixed = docx.Document(str(out))
    texts = [p.text for p in fixed.paragraphs if p.text.strip()]
    assert "• 负责产品需求文档撰写与评审推进。" in texts   # 仅符号被统一，文字不动
    assert not any(not p.text.strip() for p in fixed.paragraphs)   # 凑间距的空行被删掉
    bullet = next(p for p in fixed.paragraphs if "负责产品需求" in p.text)
    assert bullet.alignment == 0                                   # 两端对齐 → 左对齐
    assert bullet.paragraph_format.line_spacing == profile["spacing"]["line_spacing"]
    assert round(bullet.paragraph_format.first_line_indent.cm, 1) == -profile["bullet"]["hanging_cm"]
    assert round(fixed.sections[0].left_margin.cm, 2) == profile["page"]["margin_left_cm"]


def test_format_docx_reports_header_contact(bad_docx: Path) -> None:
    warnings = format_docx.structural_warnings(docx.Document(str(bad_docx)))
    assert any("页眉" in w and w.startswith("🔴") for w in warnings)


def test_format_docx_refuses_inplace(bad_docx: Path) -> None:
    result = subprocess.run([sys.executable, str(SCRIPTS / "format_docx.py"), str(bad_docx),
                             "-o", str(bad_docx)], capture_output=True, text=True)
    assert result.returncode != 0 and "覆盖" in result.stderr


# ---- ⑦ ATS ----------------------------------------------------------------
def test_ats_flags_bad_resume(bad_docx: Path) -> None:
    findings, _ = ats_check.check_docx(bad_docx)
    items = {f["item"]: f["level"] for f in findings}
    assert items.get("页眉") == ats_check.FATAL          # 联系方式藏在页眉
    assert items.get("邮箱") == ats_check.FATAL          # 正文里抓不到
    assert items.get("项目符号") == ats_check.HIGH       # ❖ 会变乱码


def test_ats_passes_generated_resume(parsed: dict, tmp_path: Path) -> None:
    out = tmp_path / "张三-数据分析-简历.docx"
    build_docx.render(parsed, build_docx.load_profile("zh", parsed)).save(str(out))
    findings, _ = ats_check.check_docx(out)
    fatal = [f for f in findings if f["level"] == ats_check.FATAL]
    assert not fatal, f"重建产物不应有致命项：{fatal}"


def test_ats_flags_mixed_date_formats() -> None:
    text = "工作经历 2021.03 - 2023.06 教育背景 2016年9月 至今 " + "内容" * 40
    text += " 138-1234-5678 zhangsan@example.com"
    items = {f["item"] for f in ats_check.check_text(text, ["工作经历", "教育背景"])}
    assert "日期格式" in items


def test_ats_flags_creative_headings() -> None:
    text = "我的江湖 " + "内容" * 50 + " 138-1234-5678 a@b.com 2021.03 - 2023.06"
    findings = ats_check.check_text(text, ["我的江湖"])
    assert any(f["item"] == "模块标题" for f in findings)


# ---- ⑨ 对比与零编造检查 -----------------------------------------------------
def test_diff_detects_fabricated_number_and_inflated_seniority(parsed: dict, tmp_path: Path) -> None:
    edited = json.loads(json.dumps(parsed, ensure_ascii=False))
    bullets = edited["sections"][1]["entries"][1]["bullets"]
    assert "协助" in bullets[0]["text"]
    bullets[0]["text"] = "主导日报周报体系建设，覆盖 47 个业务指标。"   # 抬职级 + 编数字
    report, issues = diff_report.render(parsed, edited)
    assert "47" in report and "编造" in report
    assert "职级" in report
    assert issues >= 2


def test_diff_allows_rewording_with_existing_numbers(parsed: dict) -> None:
    edited = json.loads(json.dumps(parsed, ensure_ascii=False))
    bullets = edited["sections"][1]["entries"][0]["bullets"]
    bullets[1]["text"] = "搭建覆盖 12 个核心指标的用户增长数据看板。"  # 只换说法，数字来自原文
    report, issues = diff_report.render(parsed, edited)
    assert issues == 0
    assert "未发现新增数字或职级抬高" in report


def test_diff_flags_leftover_placeholder(parsed: dict) -> None:
    edited = json.loads(json.dumps(parsed, ensure_ascii=False))
    edited["sections"][1]["entries"][0]["bullets"][0]["text"] = "搭建看板，覆盖【待补充：几个】指标。"
    report, issues = diff_report.render(parsed, edited)
    assert "占位符未填" in report and issues >= 1
